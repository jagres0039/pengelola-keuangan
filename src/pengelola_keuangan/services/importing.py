"""Import transactions from an Excel file produced by the bot's export."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from pengelola_keuangan.db.models import Category, Transaction, TransactionType
from pengelola_keuangan.services.categories import (
    fallback_category,
    find_category_by_name,
    get_or_create_category,
)
from pengelola_keuangan.services.time_helpers import get_zoneinfo


@dataclass(frozen=True)
class ImportRow:
    """A parsed row from an uploaded XLSX file."""

    row_index: int
    id: int | None
    occurred_at: datetime | None
    type: TransactionType
    amount: Decimal
    category_name: str
    note: str


@dataclass
class ImportPlan:
    """A preview of changes that an import would apply."""

    to_create: list[ImportRow] = field(default_factory=list)
    to_update: list[tuple[ImportRow, Transaction]] = field(default_factory=list)
    to_delete: list[Transaction] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _coerce_decimal(value: object) -> Decimal:
    """Coerce a cell value to a Decimal."""
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    if isinstance(value, str):
        return Decimal(value.replace(",", ".").strip())
    raise InvalidOperation(f"Cannot convert {value!r} to Decimal")


def _coerce_datetime(value: object, tz_name: str) -> datetime | None:
    """Coerce a cell value to a timezone-aware datetime."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.replace(tzinfo=get_zoneinfo(tz_name))
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(text, fmt)
            except ValueError:
                continue
            return parsed.replace(tzinfo=get_zoneinfo(tz_name))
    return None


def parse_import_rows(file_bytes: bytes, tz_name: str) -> tuple[list[ImportRow], list[str]]:
    """Parse an uploaded XLSX file into ImportRow objects."""
    import io

    rows: list[ImportRow] = []
    errors: list[str] = []
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = workbook["Transaksi"] if "Transaksi" in workbook.sheetnames else workbook.worksheets[0]

    raw_rows = list(sheet.iter_rows(values_only=True))
    if not raw_rows:
        errors.append("File kosong, tidak ada data.")
        return rows, errors

    header = [str(cell or "").strip().lower() for cell in raw_rows[0]]
    col = {name: idx for idx, name in enumerate(header)}
    required = {"tanggal", "tipe", "jumlah"}
    missing = required - set(col)
    if missing:
        errors.append(f"Kolom wajib hilang: {', '.join(sorted(missing))}")
        return rows, errors

    for row_idx, raw in enumerate(raw_rows[1:], start=2):
        if not raw or all(cell is None or cell == "" for cell in raw):
            continue

        def get(name: str, _raw: tuple[object, ...] = raw) -> object:
            idx = col.get(name)
            if idx is None or idx >= len(_raw):
                return None
            return _raw[idx]

        try:
            id_raw = get("id")
            id_val = int(id_raw) if isinstance(id_raw, int | float) and id_raw else None
            type_raw = str(get("tipe") or "").strip().lower()
            if type_raw not in {"in", "out"}:
                raise ValueError(f"Tipe harus 'in' atau 'out', dapat {type_raw!r}")
            tx_type = TransactionType.INCOME if type_raw == "in" else TransactionType.EXPENSE
            amount = _coerce_decimal(get("jumlah") or 0)
            if amount <= 0:
                raise ValueError(f"Jumlah harus > 0, dapat {amount}")
            occurred = _coerce_datetime(get("tanggal"), tz_name)
            category_name = str(get("kategori") or "Lainnya").strip() or "Lainnya"
            note = str(get("catatan") or "").strip()
            rows.append(
                ImportRow(
                    row_index=row_idx,
                    id=id_val,
                    occurred_at=occurred,
                    type=tx_type,
                    amount=amount,
                    category_name=category_name,
                    note=note,
                )
            )
        except (ValueError, InvalidOperation) as exc:
            errors.append(f"Baris {row_idx}: {exc}")

    return rows, errors


def build_import_plan(
    session: Session,
    user_id: int,
    rows: list[ImportRow],
    existing: list[Transaction],
) -> ImportPlan:
    """Compare uploaded rows against existing transactions and build a plan."""
    plan = ImportPlan()
    by_id = {tx.id: tx for tx in existing}
    seen_ids: set[int] = set()

    for row in rows:
        if row.id is not None and row.id in by_id:
            seen_ids.add(row.id)
            existing_tx = by_id[row.id]
            category_match = (
                existing_tx.category.name == row.category_name
                if existing_tx.category is not None
                else row.category_name in {"", "Lainnya"}
            )
            if (
                existing_tx.amount != row.amount
                or (
                    existing_tx.type.value
                    if hasattr(existing_tx.type, "value")
                    else existing_tx.type
                )
                != row.type.value
                or (existing_tx.note or "") != row.note
                or not category_match
            ):
                plan.to_update.append((row, existing_tx))
        else:
            plan.to_create.append(row)

    for tx_id, tx in by_id.items():
        if tx_id not in seen_ids:
            plan.to_delete.append(tx)

    return plan


def apply_import_plan(
    session: Session,
    user_id: int,
    plan: ImportPlan,
    tz_name: str,
) -> tuple[int, int, int]:
    """Apply an import plan to the database. Returns (created, updated, deleted)."""
    created = updated = deleted = 0

    for row in plan.to_create:
        category = _resolve_category(session, user_id, row)
        session.add(
            Transaction(
                user_id=user_id,
                type=row.type,
                amount=row.amount,
                category_id=category.id,
                note=row.note or None,
                occurred_at=row.occurred_at,
            )
        )
        created += 1

    for row, existing_tx in plan.to_update:
        category = _resolve_category(session, user_id, row)
        existing_tx.type = row.type
        existing_tx.amount = row.amount
        existing_tx.category_id = category.id
        existing_tx.note = row.note or None
        if row.occurred_at is not None:
            existing_tx.occurred_at = row.occurred_at
        updated += 1

    for tx in plan.to_delete:
        session.delete(tx)
        deleted += 1

    session.flush()
    return created, updated, deleted


def _resolve_category(session: Session, user_id: int, row: ImportRow) -> Category:
    """Find or create the right category for an import row."""
    name = row.category_name or "Lainnya"
    found = find_category_by_name(session, user_id, name, row.type)
    if found is not None:
        return found
    if name.lower() == "lainnya":
        return fallback_category(session, user_id, row.type)
    return get_or_create_category(session, user_id, name, row.type)
