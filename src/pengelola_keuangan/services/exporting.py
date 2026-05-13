"""Export transactions to CSV and XLSX files."""

from __future__ import annotations

import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from pengelola_keuangan.db.models import TransactionType
from pengelola_keuangan.services.formatting import format_money, format_month
from pengelola_keuangan.services.time_helpers import to_user_tz
from pengelola_keuangan.services.transactions import list_for_month, summarize_month


def export_month_to_csv(
    session: Session,
    user_id: int,
    year: int,
    month: int,
    tz_name: str,
) -> bytes:
    """Build a CSV with all transactions in the given month."""
    rows = list_for_month(session, user_id, year, month, tz_name)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["id", "tanggal", "tipe", "jumlah", "kategori", "catatan"])
    for transaction in rows:
        local_dt = to_user_tz(transaction.occurred_at, tz_name)
        writer.writerow(
            [
                transaction.id,
                local_dt.strftime("%Y-%m-%d %H:%M:%S"),
                transaction.type.value
                if isinstance(transaction.type, TransactionType)
                else transaction.type,
                f"{transaction.amount:.2f}",
                transaction.category.name if transaction.category else "",
                transaction.note or "",
            ]
        )
    return buf.getvalue().encode("utf-8-sig")


def export_month_to_xlsx(
    session: Session,
    user_id: int,
    year: int,
    month: int,
    tz_name: str,
    currency: str = "IDR",
) -> bytes:
    """Build an XLSX workbook with a Transactions sheet and a Summary sheet."""
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Failed to create workbook sheet.")
    sheet.title = "Transaksi"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    headers = ["ID", "Tanggal", "Tipe", "Jumlah", "Kategori", "Catatan"]
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    rows = list_for_month(session, user_id, year, month, tz_name)
    for row_idx, transaction in enumerate(rows, start=2):
        local_dt = to_user_tz(transaction.occurred_at, tz_name)
        type_value = (
            transaction.type.value
            if isinstance(transaction.type, TransactionType)
            else transaction.type
        )
        sheet.cell(row=row_idx, column=1, value=transaction.id)
        sheet.cell(row=row_idx, column=2, value=local_dt.strftime("%Y-%m-%d %H:%M:%S"))
        sheet.cell(row=row_idx, column=3, value=type_value)
        sheet.cell(row=row_idx, column=4, value=float(transaction.amount))
        sheet.cell(
            row=row_idx,
            column=5,
            value=transaction.category.name if transaction.category else "",
        )
        sheet.cell(row=row_idx, column=6, value=transaction.note or "")

    for col_idx, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = max(14, len(header) + 2)

    # Summary sheet.
    summary_sheet = workbook.create_sheet(title="Ringkasan")
    summary = summarize_month(session, user_id, year, month, tz_name)

    summary_sheet["A1"] = f"Ringkasan {format_month(year, month)}"
    summary_sheet["A1"].font = Font(bold=True, size=14)

    summary_sheet["A3"] = "Total Pemasukan"
    summary_sheet["B3"] = format_money(Decimal(summary.total_income), currency)
    summary_sheet["A4"] = "Total Pengeluaran"
    summary_sheet["B4"] = format_money(Decimal(summary.total_expense), currency)
    summary_sheet["A5"] = "Saldo Bulan"
    summary_sheet["B5"] = format_money(Decimal(summary.balance), currency)

    summary_sheet["A7"] = "Pengeluaran per Kategori"
    summary_sheet["A7"].font = Font(bold=True)
    summary_sheet["A8"] = "Kategori"
    summary_sheet["B8"] = "Total"
    summary_sheet["A8"].font = Font(bold=True)
    summary_sheet["B8"].font = Font(bold=True)
    for offset, row in enumerate(summary.expense_by_category, start=9):
        summary_sheet.cell(row=offset, column=1, value=row.category_name)
        summary_sheet.cell(row=offset, column=2, value=format_money(row.total, currency))

    base = len(summary.expense_by_category) + 11
    summary_sheet.cell(row=base, column=1, value="Pemasukan per Kategori").font = Font(bold=True)
    summary_sheet.cell(row=base + 1, column=1, value="Kategori").font = Font(bold=True)
    summary_sheet.cell(row=base + 1, column=2, value="Total").font = Font(bold=True)
    for offset, row in enumerate(summary.income_by_category, start=base + 2):
        summary_sheet.cell(row=offset, column=1, value=row.category_name)
        summary_sheet.cell(row=offset, column=2, value=format_money(row.total, currency))

    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 22

    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()
